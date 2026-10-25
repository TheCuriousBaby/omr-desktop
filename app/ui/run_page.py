"""阅卷页：选择答题卡图片、执行识别、查看与导出结果。"""
import csv
import json
import shutil
from datetime import datetime
from pathlib import Path
from typing import Callable, Dict, List, Optional

from PySide6.QtCore import Qt, QUrl
from PySide6.QtGui import QBrush, QColor, QDesktopServices, QPixmap
from PySide6.QtWidgets import (
    QCheckBox,
    QComboBox,
    QFileDialog,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QListWidget,
    QMessageBox,
    QPlainTextEdit,
    QProgressBar,
    QPushButton,
    QSplitter,
    QTableWidget,
    QTableWidgetItem,
    QToolButton,
    QVBoxLayout,
    QWidget,
)

from app.engine_client import EngineJob, write_config
from app.text_fields_io import field_names, load_texts, save_texts
from app.ui.answer_picker import AnswerPicker
from app.ui.context import AppContext
from app.ui.text_reader import TextReader
from app.ui.widgets import ImageViewer

IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".pdf"}
STATUS_TEXT = {"ok": "正常", "multi": "多涂", "error": "失败"}
STATUS_COLOR = {"multi": QColor("#fff4cc"), "error": QColor("#fde2e2")}


def collect_images(folder: Path) -> List[Path]:
    return sorted(
        p for p in folder.rglob("*") if p.is_file() and p.suffix.lower() in IMAGE_SUFFIXES
    )


def read_run_results(run_dir: Path):
    """读取一次阅卷的结果，返回 (输出列名, 记录列表)。"""
    meta = {}
    meta_path = run_dir / "run_meta.json"
    if meta_path.exists():
        meta = json.loads(meta_path.read_text("utf-8"))
    statuses: Dict[str, str] = meta.get("statuses", {})

    records: Dict[str, dict] = {}
    columns: List[str] = []
    sources = [
        (sorted((run_dir / "Results").glob("Results_*.csv")), None),
        ([run_dir / "Manual" / "MultiMarkedFiles.csv"], "multi"),
        ([run_dir / "Manual" / "ErrorFiles.csv"], "error"),
    ]
    for files, forced in sources:
        for path in files:
            if not path.exists():
                continue
            with open(path, "r", encoding="utf-8", newline="") as handle:
                reader = csv.reader(handle)
                header = next(reader, None)
                if not header:
                    continue
                if not columns:
                    columns = header[4:]
                for line in reader:
                    if len(line) < 4:
                        continue
                    name = line[0]
                    status = forced or statuses.get(name, "ok")
                    records[name] = {
                        "file": name,
                        "status": status,
                        "score": line[3],
                        "answers": line[4:],
                        "input": line[1],
                    }
    for name, status in statuses.items():
        if name not in records:
            records[name] = {
                "file": name,
                "status": status,
                "score": "",
                "answers": [""] * len(columns),
                "input": "",
            }
    for name, record in records.items():
        if name in statuses and record["status"] == "ok":
            record["status"] = statuses[name]
    return columns, sorted(records.values(), key=lambda r: r["file"])


class RunPage(QWidget):
    def __init__(self, ctx: AppContext, prepare: Callable[[], str], parent=None):
        super().__init__(parent)
        self.ctx = ctx
        self.prepare = prepare
        self.files: List[Path] = []
        self.job = EngineJob(self)
        self.job.logLine.connect(self._on_log)
        self.job.fileStarted.connect(self._on_file_started)
        self.job.fileDone.connect(self._on_file_done)
        self.job.finished.connect(self._on_finished)
        self.run_dir: Optional[Path] = None
        self.statuses: Dict[str, str] = {}
        self.columns: List[str] = []
        self.records: List[dict] = []
        self.visible_records: List[dict] = []
        self._done = 0
        self.text_fields: List[str] = []
        self.texts: Dict[str, Dict[str, str]] = {}
        self._text_run_dir: Optional[Path] = None
        self._filling = False
        self._text_note_shown = False
        self.reader = TextReader(ctx.settings, self)
        self.reader.textReady.connect(self._on_text_ready)
        self.reader.progressChanged.connect(self._on_text_progress)
        self.reader.failed.connect(lambda text: self.status_label.setText(f"大模型识别文字失败：{text}"))
        self.job.cropReady.connect(self._on_crop)
        self.setAcceptDrops(True)
        self._build_ui()
        ctx.projectChanged.connect(self.reload)

    # ---------- 界面 ----------

    def _build_ui(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(16, 16, 16, 16)

        top = QHBoxLayout()
        left = QVBoxLayout()
        row = QHBoxLayout()
        btn_folder = QPushButton("添加文件夹…")
        btn_folder.clicked.connect(self._add_folder)
        btn_files = QPushButton("添加文件…")
        btn_files.clicked.connect(self._add_files)
        btn_clear = QPushButton("清空")
        btn_clear.clicked.connect(self._clear_files)
        row.addWidget(btn_folder)
        row.addWidget(btn_files)
        row.addWidget(btn_clear)
        row.addStretch(1)
        left.addLayout(row)
        self.file_label = QLabel("可将答题卡文件或文件夹拖入此页面")
        self.file_label.setObjectName("hint")
        left.addWidget(self.file_label)
        top.addLayout(left, 1)
        self.answer_picker = AnswerPicker(self.ctx)
        top.addWidget(self.answer_picker)

        self.btn_start = QPushButton("开始阅卷")
        self.btn_start.setObjectName("primary")
        self.btn_start.setMinimumHeight(40)
        self.btn_start.clicked.connect(self.start)
        self.btn_cancel = QPushButton("取消")
        self.btn_cancel.setMinimumHeight(40)
        self.btn_cancel.clicked.connect(self.job.cancel)
        self.btn_cancel.hide()
        top.addWidget(self.btn_start)
        top.addWidget(self.btn_cancel)
        root.addLayout(top)

        self.file_list = QListWidget()
        self.file_list.setMaximumHeight(90)
        root.addWidget(self.file_list)

        progress_row = QHBoxLayout()
        self.progress = QProgressBar()
        self.progress.setTextVisible(True)
        self.progress.setFormat("%v / %m")
        self.progress.hide()
        self.status_label = QLabel("")
        progress_row.addWidget(self.progress, 1)
        progress_row.addWidget(self.status_label)
        root.addLayout(progress_row)

        tool = QHBoxLayout()
        tool.addWidget(QLabel("历史记录"))
        self.history_combo = QComboBox()
        self.history_combo.setMinimumWidth(240)
        self.history_combo.activated.connect(self._on_history_selected)
        tool.addWidget(self.history_combo)
        tool.addSpacing(12)
        tool.addWidget(QLabel("筛选"))
        self.filter_combo = QComboBox()
        for key, title in [("all", "全部"), ("abnormal", "仅异常（多涂、失败）"), ("blank", "含未涂题")]:
            self.filter_combo.addItem(title, key)
        self.filter_combo.currentIndexChanged.connect(self._apply_filter)
        tool.addWidget(self.filter_combo)
        tool.addStretch(1)
        self.log_toggle = QCheckBox("显示日志")
        self.log_toggle.toggled.connect(lambda on: self.log_view.setVisible(on))
        tool.addWidget(self.log_toggle)
        self.btn_open = QPushButton("打开输出目录")
        self.btn_open.clicked.connect(self._open_output)
        self.btn_export = QPushButton("导出结果…")
        self.btn_export.clicked.connect(self._export)
        tool.addWidget(self.btn_open)
        tool.addWidget(self.btn_export)
        root.addLayout(tool)

        splitter = QSplitter(Qt.Horizontal)
        self.table = QTableWidget(0, 0)
        self.table.setEditTriggers(QTableWidget.DoubleClicked | QTableWidget.EditKeyPressed)
        self.table.setSelectionBehavior(QTableWidget.SelectRows)
        self.table.setSelectionMode(QTableWidget.SingleSelection)
        self.table.verticalHeader().hide()
        self.table.currentCellChanged.connect(self._on_row_selected)
        self.table.itemChanged.connect(self._on_cell_edited)
        splitter.addWidget(self.table)

        viewer_box = QWidget()
        viewer_layout = QVBoxLayout(viewer_box)
        viewer_layout.setContentsMargins(0, 0, 0, 0)
        viewer_bar = QHBoxLayout()
        self.viewer_title = QLabel("标注图")
        viewer_bar.addWidget(self.viewer_title, 1)
        for text, callback in [("适应", lambda: self.viewer.fit()), ("放大", lambda: self.viewer.zoom(1.25)), ("缩小", lambda: self.viewer.zoom(0.8))]:
            button = QToolButton()
            button.setText(text)
            button.clicked.connect(callback)
            viewer_bar.addWidget(button)
        self.btn_original = QPushButton("打开原图")
        self.btn_original.clicked.connect(self._open_original)
        viewer_bar.addWidget(self.btn_original)
        viewer_layout.addLayout(viewer_bar)
        self.viewer = ImageViewer()
        viewer_layout.addWidget(self.viewer, 1)
        splitter.addWidget(viewer_box)
        splitter.setStretchFactor(0, 3)
        splitter.setStretchFactor(1, 2)
        root.addWidget(splitter, 1)

        self.log_view = QPlainTextEdit()
        self.log_view.setReadOnly(True)
        self.log_view.setMaximumHeight(150)
        self.log_view.setMaximumBlockCount(5000)
        self.log_view.hide()
        root.addWidget(self.log_view)

    # ---------- 输入文件 ----------

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()

    def dropEvent(self, event):
        paths = [Path(url.toLocalFile()) for url in event.mimeData().urls() if url.isLocalFile()]
        self._add_paths(paths)

    def _add_folder(self):
        folder = QFileDialog.getExistingDirectory(self, "选择答题卡所在文件夹")
        if folder:
            self._add_paths([Path(folder)])

    def _add_files(self):
        files, _ = QFileDialog.getOpenFileNames(self, "选择答题卡", "", "答题卡 (*.jpg *.jpeg *.png *.pdf)")
        self._add_paths([Path(f) for f in files])

    def _add_paths(self, paths: List[Path]):
        known = set(self.files)
        for path in paths:
            found = collect_images(path) if path.is_dir() else [path]
            for item in found:
                if item.suffix.lower() in IMAGE_SUFFIXES and item not in known:
                    known.add(item)
                    self.files.append(item)
        self._refresh_files()

    def _clear_files(self):
        self.files = []
        self._refresh_files()

    def _refresh_files(self):
        self.file_list.clear()
        for path in self.files:
            self.file_list.addItem(str(path))
        count = len(self.files)
        self.file_label.setText(f"已选择 {count} 个文件" if count else "可将答题卡文件或文件夹拖入此页面")

    # ---------- 执行 ----------

    @property
    def project(self):
        return self.ctx.project

    def reload(self):
        self.setEnabled(self.project is not None)
        self._refresh_history()
        self._show_records([], [])
        self.viewer.show_text("")
        if self.history_combo.count():
            self.history_combo.setCurrentIndex(0)
            self._on_history_selected(0)

    def start(self):
        project = self.project
        if not project or self.job.running():
            return
        if not self.files:
            QMessageBox.information(self, "阅卷", "请先添加答题卡图片或文件夹。")
            return
        error = self.prepare()
        if error:
            QMessageBox.warning(self, "无法开始", error)
            return

        output_root = Path(self.ctx.settings.get("app.output_root"))
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        self.run_dir = output_root / project.name / stamp
        run_meta_dir = self.run_dir / "_run"
        run_meta_dir.mkdir(parents=True, exist_ok=True)
        config_path = run_meta_dir / "config.json"
        write_config(config_path, project.engine_config(self.ctx.settings))
        files_path = run_meta_dir / "files.txt"
        files_path.write_text("\n".join(str(p) for p in self.files), "utf-8")
        snapshot = run_meta_dir / "project"
        snapshot.mkdir(exist_ok=True)
        shutil.copyfile(project.dir / "template.json", snapshot / "template.json")
        chosen_answer = self.answer_picker.selected_path()
        if chosen_answer:
            shutil.copyfile(chosen_answer, snapshot / "evaluation.json")

        fields = project.text_fields
        self.text_fields = [f["name"] for f in fields]
        self.texts = {}
        self._text_run_dir = self.run_dir
        self._text_note_shown = False
        self.reader.reset()
        engine_options = dict(self.answer_picker.engine_options())
        if fields:
            textfields_path = run_meta_dir / "textfields.json"
            project.write_text_fields(textfields_path)
            engine_options["textfields"] = textfields_path

        self.statuses = {}
        self._done = 0
        self.log_view.clear()
        self.progress.setRange(0, len(self.files))
        self.progress.setValue(0)
        self.progress.show()
        self.status_label.setText("正在识别…")
        self._show_records([], [])
        self.viewer.show_text("")
        self._set_running(True)
        self.job.start(
            "run",
            project=project.dir,
            config=config_path,
            files=files_path,
            output=self.run_dir,
            **engine_options,
        )

    def _set_running(self, running: bool):
        self.btn_start.setVisible(not running)
        self.btn_cancel.setVisible(running)
        self.history_combo.setEnabled(not running)

    def _on_log(self, line: str):
        self.log_view.appendPlainText(line)

    def _on_file_started(self, name: str):
        self.status_label.setText(f"正在识别：{name}")

    def _on_file_done(self, name: str, status: str):
        self.statuses[name] = status
        self._done += 1
        self.progress.setMaximum(max(self.progress.maximum(), self._done))
        self.progress.setValue(self._done)

    def _on_finished(self, result: dict):
        self._set_running(False)
        meta = {
            "project": self.project.name if self.project else "",
            "finished": datetime.now().isoformat(timespec="seconds"),
            "total": len(self.statuses),
            "statuses": self.statuses,
            "inputs": [str(p) for p in self.files],
            "ok": bool(result.get("ok")),
        }
        if self.run_dir and self.run_dir.exists():
            (self.run_dir / "run_meta.json").write_text(
                json.dumps(meta, ensure_ascii=False, indent=2), "utf-8"
            )
        if result.get("cancelled"):
            self.status_label.setText("已取消")
        elif not result.get("ok"):
            self.status_label.setText("识别失败")
            self.log_toggle.setChecked(True)
            QMessageBox.warning(self, "识别失败", result.get("error", "未知错误"))
        else:
            abnormal = sum(1 for s in self.statuses.values() if s != "ok")
            self.status_label.setText(
                f"完成：共 {len(self.statuses)} 份，异常 {abnormal} 份"
            )
        self._refresh_history()
        if self.run_dir and (self.run_dir / "run_meta.json").exists():
            self._load_run(self.run_dir)
            index = self.history_combo.findData(str(self.run_dir))
            if index >= 0:
                self.history_combo.setCurrentIndex(index)
            if result.get("ok") and self.ctx.settings.get("app.open_after_run"):
                self._open_output()

    # ---------- 文字区域（姓名等）识别 ----------

    def _on_crop(self, info: dict):
        if not self.reader.available():
            if not self._text_note_shown:
                self._text_note_shown = True
                self.status_label.setText("已定义文字区域，但未在设置中配置大模型，文字没有识别。")
            return
        self.reader.enqueue(info["file"], info["field"], info["path"])

    def _on_text_ready(self, file: str, field: str, text: str):
        self.texts.setdefault(file, {})[field] = text
        if self._text_run_dir is not None:
            save_texts(self._text_run_dir, self.text_fields, self.texts)
        if field not in self.text_fields:
            return
        column = 3 + self.text_fields.index(field)
        for row, record in enumerate(self.visible_records):
            if record["file"] == file and self.table.columnCount() > column:
                self._filling = True
                item = self.table.item(row, column)
                if item is None:
                    item = QTableWidgetItem()
                    self.table.setItem(row, column, item)
                item.setText(text)
                self._filling = False
                break

    def _on_text_progress(self, done: int, total: int):
        if total and not self.job.running():
            self.status_label.setText(
                "文字识别完成" if done >= total else f"正在识别文字 {done}/{total}…"
            )

    def _on_cell_edited(self, item):
        if self._filling or not self.text_fields:
            return
        column = item.column() - 3
        row = item.row()
        if not (0 <= column < len(self.text_fields) and 0 <= row < len(self.visible_records)):
            return
        file = self.visible_records[row]["file"]
        self.texts.setdefault(file, {})[self.text_fields[column]] = item.text().strip()
        if self.run_dir is not None:
            save_texts(self.run_dir, self.text_fields, self.texts)

    # ---------- 结果展示 ----------

    def _refresh_history(self):
        self.history_combo.clear()
        project = self.project
        if not project:
            return
        root = Path(self.ctx.settings.get("app.output_root")) / project.name
        if not root.exists():
            return
        for directory in sorted((d for d in root.iterdir() if d.is_dir()), reverse=True):
            meta_path = directory / "run_meta.json"
            if not meta_path.exists():
                continue
            try:
                meta = json.loads(meta_path.read_text("utf-8"))
                moment = datetime.strptime(directory.name, "%Y%m%d_%H%M%S")
                title = f"{moment:%Y-%m-%d %H:%M:%S}（{meta.get('total', 0)} 份）"
            except ValueError:
                title = directory.name
            self.history_combo.addItem(title, str(directory))

    def _on_history_selected(self, index):
        path = self.history_combo.itemData(index)
        if path:
            self._load_run(Path(path))

    def _load_run(self, run_dir: Path):
        self.run_dir = run_dir
        columns, records = read_run_results(run_dir)
        self.text_fields = field_names(run_dir)
        loaded = load_texts(run_dir)
        if run_dir == self._text_run_dir:
            for file, values in self.texts.items():
                loaded.setdefault(file, {}).update(values)
        self.texts = loaded
        self._show_records(columns, records)

    def _show_records(self, columns: List[str], records: List[dict]):
        self.columns = columns
        self.records = records
        self._apply_filter()

    def _apply_filter(self):
        mode = self.filter_combo.currentData()
        rows = []
        for record in self.records:
            if mode == "abnormal" and record["status"] == "ok":
                continue
            if mode == "blank" and not any(a == "" for a in record["answers"]):
                continue
            rows.append(record)
        self.visible_records = rows
        headers = ["文件", "状态", "得分"] + self.text_fields + self.columns
        text_start, text_end = 3, 3 + len(self.text_fields)
        self._filling = True
        self.table.setColumnCount(len(headers))
        self.table.setHorizontalHeaderLabels(headers)
        self.table.setRowCount(len(rows))
        for r, record in enumerate(rows):
            file_texts = self.texts.get(record["file"], {})
            values = (
                [record["file"], STATUS_TEXT.get(record["status"], record["status"]), record["score"]]
                + [file_texts.get(name, "") for name in self.text_fields]
                + record["answers"]
            )
            for c, value in enumerate(values):
                item = QTableWidgetItem(value)
                color = STATUS_COLOR.get(record["status"])
                if color:
                    item.setBackground(QBrush(color))
                if not text_start <= c < text_end:
                    item.setFlags(item.flags() & ~Qt.ItemIsEditable)
                self.table.setItem(r, c, item)
        self._filling = False
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        self.table.horizontalHeader().setStretchLastSection(False)

    def _selected_record(self) -> Optional[dict]:
        row = self.table.currentRow()
        if 0 <= row < len(self.visible_records):
            return self.visible_records[row]
        return None

    def _on_row_selected(self, *_):
        record = self._selected_record()
        if not record or not self.run_dir:
            return
        self.viewer_title.setText(f"标注图：{record['file']}")
        marked = self.run_dir / "CheckedOMRs" / record["file"]
        if marked.exists():
            pixmap = QPixmap(str(marked))
            self.viewer.set_pixmap(pixmap)
        else:
            self.viewer.show_text("没有标注图（设置中关闭了“保存标注图”，或该文件识别失败）")

    def _open_original(self):
        record = self._selected_record()
        if not record:
            return
        candidates = [Path(record["input"])] if record["input"] else []
        meta_path = self.run_dir / "run_meta.json" if self.run_dir else None
        if meta_path and meta_path.exists():
            inputs = json.loads(meta_path.read_text("utf-8")).get("inputs", [])
            candidates += [Path(p) for p in inputs if Path(p).stem == Path(record["file"]).stem]
        for path in candidates:
            if path.exists():
                QDesktopServices.openUrl(QUrl.fromLocalFile(str(path)))
                return
        QMessageBox.information(self, "原图", "找不到原图文件。")

    def _open_output(self):
        if self.run_dir and self.run_dir.exists():
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(self.run_dir)))

    # ---------- 导出 ----------

    def _export(self):
        if not self.records:
            QMessageBox.information(self, "导出", "没有可导出的结果。")
            return
        default = f"{self.project.name}_结果.xlsx" if self.project else "结果.xlsx"
        path, selected = QFileDialog.getSaveFileName(
            self, "导出结果", default, "Excel 工作簿 (*.xlsx);;CSV 文件 (*.csv)"
        )
        if not path:
            return
        headers = ["文件", "状态", "得分"] + self.text_fields + self.columns
        rows = [
            [r["file"], STATUS_TEXT.get(r["status"], r["status"]), r["score"]]
            + [self.texts.get(r["file"], {}).get(name, "") for name in self.text_fields]
            + r["answers"]
            for r in self.records
        ]
        try:
            if path.lower().endswith(".csv"):
                with open(path, "w", encoding="utf-8-sig", newline="") as handle:
                    writer = csv.writer(handle)
                    writer.writerow(headers)
                    writer.writerows(rows)
            else:
                from openpyxl import Workbook
                from openpyxl.styles import PatternFill

                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "结果"
                sheet.append(headers)
                fills = {
                    "多涂": PatternFill("solid", fgColor="FFF4CC"),
                    "失败": PatternFill("solid", fgColor="FDE2E2"),
                }
                for row in rows:
                    sheet.append(row)
                    fill = fills.get(row[1])
                    if fill:
                        for cell in sheet[sheet.max_row]:
                            cell.fill = fill
                workbook.save(path)
        except OSError as error:
            QMessageBox.warning(self, "导出失败", str(error))
            return
        QMessageBox.information(self, "导出", "导出完成。")
